import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_styled_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    h.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    h.paragraph_format.space_after = Pt(4)
    run = h.runs[0]
    run.font.name = 'Calibri'
    if level == 1:
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = RGBColor(16, 44, 87) # Deep Navy
    elif level == 2:
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 86, 160) # Royal Blue
    elif level == 3:
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(45, 55, 72) # Slate
    return h

def format_inline_text(paragraph, text):
    # Splits inline code `code`, bold **text**, italics *text*, and links [text](url)
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\[.*?\]\(.*?\))', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(180, 40, 40)
        elif token.startswith('[') and '](' in token and token.endswith(')'):
            m = re.match(r'\[(.*?)\]\((.*?)\)', token)
            if m:
                txt, url = m.groups()
                run = paragraph.add_run(f"{txt} ({url})")
                run.font.color.rgb = RGBColor(0, 102, 204)
                run.underline = True
            else:
                paragraph.add_run(token)
        else:
            paragraph.add_run(token)

def convert_markdown_to_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = docx.Document()
    
    # Page setup - standard 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Set base styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(40, 40, 40)

    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    def flush_table(tbl_lines):
        if not tbl_lines:
            return
        parsed_rows = []
        for line in tbl_lines:
            line_str = line.strip()
            if not line_str or line_str.startswith('|---') or line_str.startswith('| ---'):
                continue
            cols = [c.strip() for c in line_str.strip('|').split('|')]
            if cols:
                parsed_rows.append(cols)
        
        if not parsed_rows:
            return
        
        num_cols = max(len(r) for r in parsed_rows)
        table = doc.add_table(rows=len(parsed_rows), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        
        for r_idx, row in enumerate(parsed_rows):
            for c_idx in range(num_cols):
                cell = table.cell(r_idx, c_idx)
                val = row[c_idx] if c_idx < len(row) else ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                format_inline_text(p, val)
                
                # Header styling
                if r_idx == 0:
                    set_cell_background(cell, "1E3A8A") # Navy blue
                    for run in p.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                else:
                    bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
                    set_cell_background(cell, bg)
                set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
        
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def flush_code(lines_to_code):
        if not lines_to_code:
            return
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.left_indent = Inches(0.2)
        code_text = "\n".join(lines_to_code)
        run = p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(33, 37, 41)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Code block fence
        if stripped.startswith('```'):
            if in_code_block:
                in_code_block = False
                flush_code(code_lines)
                code_lines = []
            else:
                if in_table:
                    flush_table(table_lines)
                    table_lines = []
                    in_table = False
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line.rstrip('\r\n'))
            i += 1
            continue

        # Table detection
        if stripped.startswith('|') and stripped.endswith('|'):
            if not in_table:
                in_table = True
                table_lines = []
            table_lines.append(stripped)
            i += 1
            continue
        else:
            if in_table:
                flush_table(table_lines)
                table_lines = []
                in_table = False

        # Blank line
        if not stripped:
            i += 1
            continue

        # Headings
        if stripped.startswith('# '):
            add_styled_heading(doc, stripped[2:], level=1)
        elif stripped.startswith('## '):
            add_styled_heading(doc, stripped[3:], level=2)
        elif stripped.startswith('### '):
            add_styled_heading(doc, stripped[4:], level=3)
        elif stripped.startswith('#### '):
            add_styled_heading(doc, stripped[5:], level=4)
        elif stripped.startswith('---'):
            # Divider
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run("―" * 45)
            run.font.color.rgb = RGBColor(200, 200, 200)
        elif stripped.startswith('* ') or stripped.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            format_inline_text(p, stripped[2:])
        elif re.match(r'^\d+\.\s', stripped):
            num_match = re.match(r'^(\d+\.)\s(.*)$', stripped)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            format_inline_text(p, num_match.group(2))
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(5)
            format_inline_text(p, stripped)

        i += 1

    if in_table:
        flush_table(table_lines)
    if in_code_block:
        flush_code(code_lines)

    try:
        doc.save(docx_path)
        print(f"Successfully generated DOCX at {docx_path}")
    except PermissionError:
        alt_path = docx_path.replace(".docx", "_Updated.docx")
        doc.save(alt_path)
        print(f"Notice: {docx_path} is currently locked (open in Word). Successfully saved updated version to {alt_path}")

if __name__ == '__main__':
    md_file = r"c:\Users\srava\Downloads\Ridgevision-ai-main\Ridgevision-ai-main\RidgeVision_v2_Comprehensive_Master_Report.md"
    docx_file = r"c:\Users\srava\Downloads\Ridgevision-ai-main\Ridgevision-ai-main\RidgeVision_v2_Comprehensive_Master_Report.docx"
    convert_markdown_to_docx(md_file, docx_file)
